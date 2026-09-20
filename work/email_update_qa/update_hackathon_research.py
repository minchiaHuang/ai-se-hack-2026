from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

FILES = {
    "AI_for_Social_Enterprise_Hackathon_Research_Brief_v1.2.docx": {
        "date": "研究日期：2026 年 9 月 17 日\n活動日期：2026 年 9 月 21–22 日\n團隊：Developer × Designer × Data Science",
        "intro": "本文件整合截至 2026 年 9 月 17 日的公開資料、主辦方 9 月 16 日 email 與 LinkedIn 唯讀調查。內容分為已確認事實、研究推論與未確認事項。主辦方及評審背景屬於公開紀錄推論，不能視為正式評分標準。",
        "time": "2026/9/21（一）8:30 AM 報到；9:00 AM 至 2026/9/22（二）5:00 PM，AEST",
        "day1": "8:30 報到；10:15 問題提案；10:45 組隊並開始建置；13:00–15:00 非技術導師；15:00–17:00 技術導師；17:00–20:30 Spark 開幕活動（可選）",
        "prize": "獲勝隊伍每位成員可獲 Goods 4 Good prize pack，並新增 City Quokka 贊助的《The Innovator's Playbook》一冊。其他獎項仍待公布。",
        "prep": "先登入至少一個 AI assistant；若打算開發，先登入一個 prototype 工具，或確認自己的 API 存取方式。",
        "pitch": "若想在 Day 1 提題，準備 2 分鐘說明：服務誰、問題與代價、為何 AI 可能有幫助，以及想找哪些技能。這是可選項，不需 solution、簡報或既有團隊。",
        "social": "Social Traders 的五項選題檢查：目的（解決社會／環境問題）、營運（優先使命、人與環境）、收入（可持續）、盈餘用途（多數回投使命）、結構（以法律與財務安排長期鎖定使命）。這是選題與落地框架，不是本次正式評分 rubric 或參賽資格。",
        "designer": "前後對比流程、可理解介面、2 分鐘問題提案；Day 2 最終 pitch 時長待確認",
        "source": "主辦方 email（2026 年 9 月 16 日）｜活動已額滿、完整 Day 1 時間表與行前準備\nSocial Traders｜What is a social enterprise? https://www.socialtraders.com.au/what-is-a-social-enterprise",
        "status": "文件狀態：研究、email 已確認活動資訊、評審公開訊號、澳洲問題資料與完整訪談領域更新版 v1.5｜截至 2026 年 9 月 17 日",
    },
    "AI_for_Social_Enterprise_Hackathon_Research_Brief_v1.4_Simplified_Chinese.docx": {
        "date": "研究日期：2026 年 9 月 17 日\n活动日期：2026 年 9 月 21–22 日\n团队：Developer × Designer × Data Science",
        "intro": "本文件整合截至 2026 年 9 月 17 日的公开资料、主办方 9 月 16 日 email 与 LinkedIn 唯读调查。内容分为已确认事实、研究推论与未确认事项。主办方及评审背景属于公开纪录推论，不能视为正式评分标准。",
        "time": "2026/9/21（一）8:30 AM 报到；9:00 AM 至 2026/9/22（二）5:00 PM，AEST",
        "day1": "8:30 报到；10:15 问题提案；10:45 组队并开始建置；13:00–15:00 非技术导师；15:00–17:00 技术导师；17:00–20:30 Spark 开幕活动（可选）",
        "prize": "获胜队伍每位成员可获 Goods 4 Good prize pack，并新增 City Quokka 赞助的《The Innovator's Playbook》一册。其他奖项仍待公布。",
        "prep": "先登入至少一个 AI assistant；若打算开发，先登入一个 prototype 工具，或确认自己的 API 存取方式。",
        "pitch": "若想在 Day 1 提题，准备 2 分钟说明：服务谁、问题与代价、为何 AI 可能有帮助，以及想找哪些技能。这是可选项，不需 solution、简报或既有团队。",
        "social": "Social Traders 的五项选题检查：目的（解决社会／环境问题）、营运（优先使命、人和环境）、收入（可持续）、盈余用途（多数回投使命）、结构（以法律与财务安排长期锁定使命）。这是选题与落地框架，不是本次正式评分 rubric 或参赛资格。",
        "designer": "前后对比流程、可理解介面、2 分钟问题提案；Day 2 最终 pitch 时长待确认",
        "source": "主办方 email（2026 年 9 月 16 日）｜活动已额满、完整 Day 1 时间表与行前准备\nSocial Traders｜What is a social enterprise? https://www.socialtraders.com.au/what-is-a-social-enterprise",
        "status": "文件状态：研究、email 已确认活动资讯、评审公开讯号、澳洲问题资料与完整访谈领域更新版 v1.5｜截至 2026 年 9 月 17 日",
    },
    "AI_for_Social_Enterprise_Hackathon_Research_Brief_v1.4_Korean.docx": {
        "date": "연구일자: 2026년 9월 17일\n이벤트 기간: 2026년 9월 21-22일\n팀: Developer × Designer × Data Science",
        "intro": "이 문서는 2026년 9월 17일까지의 공개 자료, 주최 측의 9월 16일 이메일 및 LinkedIn 읽기 전용 조사를 통합합니다. 내용은 확인된 사실, 연구 추론 및 미확인 사항으로 구분하며, 주최자와 심사위원의 배경은 공식 평가 기준이 아닙니다.",
        "time": "2026/9/21(월) 8:30 AM 체크인; 9:00 AM부터 2026/9/22(화) 5:00 PM까지, AEST",
        "day1": "8:30 체크인; 10:15 문제 제안; 10:45 팀 구성 및 시작; 13:00-15:00 비기술 멘토링; 15:00-17:00 기술 멘토링; 17:00-20:30 Spark 개막 행사(선택)",
        "prize": "우승팀 각 구성원은 Goods 4 Good prize pack과 City Quokka가 후원하는 The Innovator's Playbook 한 권을 받습니다. 추가 상품은 추후 공지됩니다.",
        "prep": "최소 한 개의 AI assistant에 미리 로그인하세요. 개발할 경우 prototype 도구 하나에 로그인하거나 자신의 API 접근 방식을 확인하세요.",
        "pitch": "Day 1에 문제를 제안하고 싶다면 2분 동안 대상 사용자, 문제와 비용, AI가 도움이 될 수 있는 이유, 필요한 팀 기술을 설명하세요. 선택 사항이며 solution, 슬라이드, 기존 팀은 필요하지 않습니다.",
        "social": "Social Traders의 다섯 가지 선택 기준: 목적(사회 또는 환경 문제 해결), 운영(사명·사람·환경 우선), 수익(지속 가능), 잉여금 사용(대부분 사명에 재투자), 구조(법률·재무 구조로 장기 목적 보호). 이는 문제 선택과 도입을 위한 틀이며 이번 대회의 공식 채점 기준이나 참가 자격은 아닙니다.",
        "designer": "전후 비교 흐름, 이해하기 쉬운 인터페이스, 2분 문제 제안; Day 2 최종 pitch 시간은 미확인",
        "source": "주최 측 이메일(2026년 9월 16일) | 행사 매진, Day 1 전체 일정 및 사전 준비\nSocial Traders | What is a social enterprise? https://www.socialtraders.com.au/what-is-a-social-enterprise",
        "status": "파일 상태: 연구, 이메일로 확인한 행사 정보, 심사 공개 신호, 호주 문제 자료 및 인터뷰 분야 업데이트 v1.5 | 2026년 9월 17일 기준",
    },
}

def set_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)

def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new = Paragraph(new_p, paragraph._parent)
    new.style = paragraph.style
    if paragraph.runs:
        new_run = new.add_run(text)
        if paragraph.runs[0]._element.rPr is not None:
            new_run._element.insert(0, deepcopy(paragraph.runs[0]._element.rPr))
    else:
        new.add_run(text)
    return new

def find_paragraph(doc, contains):
    for p in doc.paragraphs:
        if contains in p.text:
            return p
    raise RuntimeError(f"Could not find paragraph containing {contains!r}")

for path, t in FILES.items():
    doc = Document(path)
    set_text(doc.paragraphs[5], t["date"])
    set_text(doc.paragraphs[8 if "Simplified" not in path and "Korean" not in path else (8 if "Simplified" in path else 7)], t["intro"])
    set_text(doc.tables[0].cell(2, 1).paragraphs[0], t["time"])
    set_text(doc.tables[0].cell(5, 1).paragraphs[0], t["day1"])
    set_text(doc.paragraphs[21 if "Korean" not in path else 20], t["prize"])
    prep_anchor = doc.paragraphs[22 if "Korean" not in path else 21]
    prep_p = insert_after(prep_anchor, t["prep"])
    pitch_p = insert_after(prep_p, t["pitch"])
    insert_after(pitch_p, t["social"])
    set_text(doc.tables[8].cell(2, 2).paragraphs[0], t["designer"])
    source_anchor = find_paragraph(doc, "Social Traders｜Sustainable Impact" if "Korean" not in path else "Social Traders｜Sustainable Impact")
    insert_after(source_anchor, t["source"])
    status_anchor = find_paragraph(doc, "文件狀態" if "v1.2" in path else ("文件状态" if "Simplified" in path else "파일 상태"))
    set_text(status_anchor, t["status"])
    doc.save(path)
    print(f"updated {path}")
