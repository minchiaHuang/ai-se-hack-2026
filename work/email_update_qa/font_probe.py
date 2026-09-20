from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

FONTS = ["Arial Unicode MS", "Songti TC", "Heiti TC", "AppleGothic", "Hiragino Sans GB", "Apple SD Gothic Neo"]
doc = Document()
for font in FONTS:
    p = doc.add_paragraph()
    run = p.add_run(f"{font}: 繁體中文 简体中文 한국어")
    run.font.size = Pt(18)
    run.font.name = font
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), font)
doc.save("work/email_update_qa/font_probe.docx")
