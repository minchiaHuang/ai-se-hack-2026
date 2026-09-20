import re
from docx import Document


PATH = "AI_for_Social_Enterprise_Hackathon_Research_Brief_v1.4_Korean.docx"
PLACEHOLDER = re.compile(r"\s*\{+\d*\}+")


def clean_paragraph(paragraph):
    for run in paragraph.runs:
        cleaned = PLACEHOLDER.sub("", run.text)
        if cleaned != run.text:
            run.text = cleaned


def walk_table(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                clean_paragraph(paragraph)
            for nested in cell.tables:
                walk_table(nested)


doc = Document(PATH)
for paragraph in doc.paragraphs:
    clean_paragraph(paragraph)
for table in doc.tables:
    walk_table(table)
for section in doc.sections:
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        clean_paragraph(paragraph)
doc.save(PATH)
